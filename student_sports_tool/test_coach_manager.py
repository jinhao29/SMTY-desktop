# -*- coding: utf-8 -*-
"""教练档案数据层单测：ensure_file / save_coach upsert / 软删除离职与恢复。"""
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import coach_manager as cm


@pytest.fixture()
def archive_dir(tmp_path):
    return str(tmp_path / 'archives')


def test_ensure_file_creates_header(archive_dir):
    fpath = cm.ensure_file(archive_dir)
    assert os.path.exists(fpath)
    from openpyxl import load_workbook
    wb = load_workbook(fpath)
    ws = wb[cm.COACH_SHEET]
    got = [ws.cell(row=1, column=i).value for i in range(1, len(cm.HEADERS) + 1)]
    assert got == cm.HEADERS
    assert cm.META_SHEET in wb.sheetnames


def test_save_and_list_coach(archive_dir):
    cm.save_coach(archive_dir, {
        'name': '王教练', 'phone': '13800000000', 'role': '主教练',
        'specialty': '青少年体能', 'join_date': '2025-09-01', 'note': '',
    })
    coaches = cm.list_coaches(archive_dir)
    assert len(coaches) == 1
    c = coaches[0]
    assert c['name'] == '王教练'
    assert c['phone'] == '13800000000'
    assert c['role'] == '主教练'
    assert c['specialty'] == '青少年体能'
    assert c['join_date'] == '2025-09-01'
    assert c['is_active'] is True


def test_save_coach_upsert_by_name(archive_dir):
    cm.save_coach(archive_dir, {'name': '李教练', 'phone': '111'})
    cm.save_coach(archive_dir, {'name': '李教练', 'phone': '222', 'role': '体能教练'})
    coaches = cm.list_coaches(archive_dir)
    assert len(coaches) == 1
    assert coaches[0]['phone'] == '222'
    assert coaches[0]['role'] == '体能教练'


def test_save_coach_empty_name_raises(archive_dir):
    with pytest.raises(ValueError):
        cm.save_coach(archive_dir, {'name': '  '})


def test_soft_delete_and_reactivate(archive_dir):
    cm.save_coach(archive_dir, {'name': '张教练', 'phone': '333'})
    # 离职（软删除）：默认列表不可见
    assert cm.set_coach_active(archive_dir, '张教练', False) is True
    assert cm.list_coaches(archive_dir) == []
    inactive = cm.list_coaches(archive_dir, include_inactive=True)
    assert len(inactive) == 1 and inactive[0]['is_active'] is False
    # 恢复在职
    assert cm.set_coach_active(archive_dir, '张教练', True) is True
    assert len(cm.list_coaches(archive_dir)) == 1


def test_set_active_missing_coach_returns_false(archive_dir):
    assert cm.set_coach_active(archive_dir, '不存在', False) is False


def test_multiple_coaches_listing(archive_dir):
    for i, name in enumerate(['教练甲', '教练乙', '教练丙']):
        cm.save_coach(archive_dir, {'name': name, 'phone': str(i)})
    cm.set_coach_active(archive_dir, '教练丙', False)
    active = cm.list_coaches(archive_dir)
    assert [c['name'] for c in active] == ['教练甲', '教练乙']
    allc = cm.list_coaches(archive_dir, include_inactive=True)
    assert len(allc) == 3
