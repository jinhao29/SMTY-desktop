# -*- coding: utf-8 -*-
"""跨端一致性测试：PC(standards.py / scorer.py) ↔ Android(core/Standards.kt / Scorer.kt)。

背景
----
评分标准表在双端各自维护一份，历史上已发生真实漂移：
- PC 缺 200米游泳 / 50米游泳 / 1分钟踢毽子（中考 2026 新方案实考项目）
- PC 完全没有 7-12 年级（初中/高中体测）标准
- 中考分组数量不一致（PC 一类2项/二类5项 vs Android 一类3项/二类7项）
- Android 备份导入 PC 时年级被丢弃，初高中/中考学员一律按小学生建档

本测试把 **Android 源码作为唯一真源**解析出来，与 PC 侧逐项比对，
任一端改动而未同步即失败。

发现差异时的正确做法
--------------------
先确认真源（深圳市教育局 / 国家学生体质健康标准），同步两端数值，
**不要改测试去迁就某一端**。

已知差异（暂不阻断，需后续修复，见 2026-09-12 审查报告）
-------------------------------------------------------
Android `Scorer.kt` 的负数成绩校验（v50）只覆盖非「分秒」分支：
`parseValue()` 对 unit=="分秒" 直接 return parseTime()，绕过了 `value < 0` 检查，
故分秒类项目输入 "-5" 仍会命中 `value <= full` 直接得 100 分。
PC 侧（scorer.py）已两分支同时拦截。Android 侧待修。
"""
import os
import re
import sys

import pytest

_ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _ROOT)

import standards as S  # noqa: E402
import scorer as SC  # noqa: E402

ROOT = _ROOT
KT_DIR = os.path.join(ROOT, '..', 'android_app', 'core', 'src', 'main', 'java',
                      'com', 'shangmentiyu', 'sportscoach', 'core')
KT_STANDARDS = os.path.join(KT_DIR, 'Standards.kt')
KT_SCORER = os.path.join(KT_DIR, 'Scorer.kt')

pytestmark = pytest.mark.skipif(
    not os.path.exists(KT_STANDARDS),
    reason='未找到 Android 端 Standards.kt（分发环境不含源码），跳过跨端比对')

_STD_RE = re.compile(
    r'Std\("([^"]+)",\s*"([^"]+)",\s*(MORE|LESS),\s*'
    r'([-\d.]+),\s*([-\d.]+),\s*([-\d.]+),\s*([-\d.]+)\)')
_GRADE_RE = re.compile(r'"(\d+)" to "([^"]+)"')
_NEXT_DECL_RE = re.compile(r'\n    (?:private )?val [A-Z_0-9]+')


def _read(path):
    with open(path, encoding='utf-8') as f:
        return f.read()


def _segment(src, var_name):
    """截取某个 Kotlin val 声明到下一个顶层 val 之间的源码片段。"""
    start = src.index(f'val {var_name}')
    seg = src[start:]
    m = _NEXT_DECL_RE.search(seg[10:])
    return seg[:10 + m.start()] if m else seg


def _std_tuples(chunk):
    # Kotlin 源码里方向是 MORE/LESS 常量字面量，PC 侧是小写 'more'/'less'，比对前归一化
    return [(m.group(1), m.group(2), m.group(3).lower(),
             float(m.group(4)), float(m.group(5)), float(m.group(6)), float(m.group(7)))
            for m in _STD_RE.finditer(chunk)]


def _parse_kt_block_map(src, var_name):
    """解析 Kotlin `Map<Int, List<Std>>`（PRIMARY_DATA / JUNIOR_DATA / SENIOR_DATA）。"""
    seg = _segment(src, var_name)
    marks = [(int(m.group(1)), m.start()) for m in re.finditer(r'(\d+) to listOf\(', seg)]
    assert marks, f'{var_name} 解析失败：未找到任何年级分块'
    result = {}
    for i, (grade, pos) in enumerate(marks):
        end = marks[i + 1][1] if i + 1 < len(marks) else len(seg)
        result[grade] = _std_tuples(seg[pos:end])
    return result


def _parse_kt_std_list(src, var_name):
    """解析 Kotlin `val NAME: List<Std>`（ZHONGKAO_2026）。"""
    return _std_tuples(_segment(src, var_name))


def _pc_tuples(stds):
    return [(s.name, s.unit, s.direction, float(s.boys_full), float(s.girls_full),
             float(s.boys_pass), float(s.girls_pass)) for s in stds]


@pytest.fixture(scope='module')
def kt_src():
    return _read(KT_STANDARDS)


# ---------------------------------------------------------------------------
# 1-3. 小学 / 初中 / 高中 逐年级逐项数值比对
# ---------------------------------------------------------------------------

def _assert_level_parity(kt_src, kt_var, pc_map, grades, level_name):
    kt = _parse_kt_block_map(kt_src, kt_var)
    assert set(kt) == set(grades), f'{level_name} 年级集合不一致：Android={sorted(kt)}'
    for g in grades:
        pc = _pc_tuples(pc_map.get(g, []))
        assert pc, f'PC 侧缺少 {level_name} {g} 年级标准'
        assert pc == kt[g], (
            f'{level_name} {g} 年级标准与 Android 不一致\n'
            f'  Android: {kt[g]}\n  PC     : {pc}')


@pytest.mark.parametrize('kt_var,pc_map,grades,name', [
    ('PRIMARY_DATA', S.PRIMARY, range(1, 7), '小学'),
    ('JUNIOR_DATA', S.JUNIOR, range(7, 10), '初中'),
    ('SENIOR_DATA', S.SENIOR, range(10, 13), '高中'),
])
def test_level_standards_parity(kt_src, kt_var, pc_map, grades, name):
    _assert_level_parity(kt_src, kt_var, pc_map, list(grades), name)


# ---------------------------------------------------------------------------
# 4-5. 中考标准与分组
# ---------------------------------------------------------------------------

def test_zhongkao_standards_parity(kt_src):
    kt = _parse_kt_std_list(kt_src, 'ZHONGKAO_2026')
    pc = _pc_tuples(S.ZHONGKAO_2026)
    assert pc == kt, (
        '中考标准表与 Android 不一致（顺序或数值）\n'
        f'  Android 共 {len(kt)} 项，PC 共 {len(pc)} 项\n'
        f'  Android: {[t[0] for t in kt]}\n  PC     : {[t[0] for t in pc]}')


def test_zhongkao_groups_cover_all_and_match_android_counts():
    """分组必须完整覆盖中考项目，且各组数量与 Android 注释（3/7/6）一致。"""
    all_names = [s.name for s in S.ZHONGKAO_2026]
    grouped = [n for names in S.ZHONGKAO_2026_GROUPS.values() for n in names]
    assert sorted(grouped) == sorted(all_names), (
        f'分组未完整覆盖中考项目。未分组={sorted(set(all_names) - set(grouped))}，'
        f'多余={sorted(set(grouped) - set(all_names))}')
    assert len(set(grouped)) == len(grouped), '同一项目出现在多个分组中'
    counts = [len(v) for v in S.ZHONGKAO_2026_GROUPS.values()]
    assert counts == [3, 7, 6], f'分组数量应为 一类3/二类7/三类6（与 Android 对齐），实际 {counts}'


def test_zhongkao_covers_three_previously_missing_events():
    """回归锚点：曾缺失的 3 个中考项目必须存在。"""
    names = {s.name for s in S.ZHONGKAO_2026}
    for ev in ('200米游泳', '50米游泳', '1分钟踢毽子'):
        assert ev in names, f'中考标准缺少项目：{ev}'


# ---------------------------------------------------------------------------
# 6-8. 年级编码 / 标签 / 取标准路由
# ---------------------------------------------------------------------------

def test_grade_options_parity(kt_src):
    kt = [(m.group(1), m.group(2)) for m in _GRADE_RE.finditer(_segment(kt_src, 'GRADE_OPTIONS'))]
    assert kt, 'Android GRADE_OPTIONS 解析失败'
    assert list(S.GRADE_OPTIONS) == kt, (
        f'年级编码表不一致\n  Android: {kt}\n  PC     : {list(S.GRADE_OPTIONS)}')


@pytest.mark.parametrize('code,expected', [
    ('0', ''), ('', ''), (None, ''),
    ('1', '一年级'), ('6', '六年级'),
    ('7', '初一'), ('9', '初三'),
    ('10', '高一'), ('12', '高三'),
    ('13', '中考'),
])
def test_grade_label_matches_android_rules(code, expected):
    assert S.grade_label(code) == expected


def test_grade_full_label_and_reverse_lookup():
    assert S.grade_full_label('3') == '小学三年级'
    assert S.grade_full_label('13') == '中考'
    assert S.grade_code_from_label('小学三年级') == '3'
    assert S.grade_code_from_label('初一') == '7'
    assert S.grade_code_from_label('学龄前') == '0'
    assert S.grade_code_from_label('不存在的年级') is None


@pytest.mark.parametrize('code,expect_count,expect_first', [
    ('0', 0, None),          # 学龄前无标准
    ('', 0, None),
    ('1', 4, '肺活量'),       # 小学1年级 4 项
    ('6', 6, '肺活量'),       # 小学6年级 6 项
    ('7', 6, '肺活量'),       # 初中
    ('12', 6, '肺活量'),      # 高中
])
def test_get_standards_by_grade_routing(code, expect_count, expect_first):
    stds = S.get_standards_by_grade(code)
    assert len(stds) == expect_count
    if expect_first:
        assert stds[0].name == expect_first


def test_get_standards_by_grade_returns_zhongkao_for_13():
    stds = S.get_standards_by_grade('13')
    assert stds is S.ZHONGKAO_2026
    assert len(stds) == 16


# ---------------------------------------------------------------------------
# 9. 评分公式一致性（参考实现直接抄自 Android Scorer.kt:92-111）
# ---------------------------------------------------------------------------

def _kt_score_more(value, full, pass_):
    """逐行对应 Scorer.kt scoreMore。"""
    if value >= full:
        return 100.0
    if value <= pass_:
        if pass_ == 0.0:
            return 0.0 if value <= 0 else max(0.0, value / full * 60)
        return max(0.0, value / pass_) * 60.0
    return 60.0 + (value - pass_) / (full - pass_) * 40.0


def _kt_score_less(value, full, pass_):
    """逐行对应 Scorer.kt scoreLess。"""
    if value <= full:
        return 100.0
    if value >= pass_:
        extra = value - pass_
        span = pass_ - full
        if span <= 0:
            return 30.0
        return max(0.0, 60.0 - extra / span * 60.0)
    return 60.0 + (pass_ - value) / (pass_ - full) * 40.0


@pytest.mark.parametrize('direction,value,full,pass_', [
    ('more', 1700, 1700, 700),   # 达到满分
    ('more', 700, 1700, 700),    # 恰好及格
    ('more', 1200, 1700, 700),   # 区间内插值
    ('more', 350, 1700, 700),    # 低于及格一半
    ('more', 0, 1700, 700),      # 零值
    ('less', 245, 245, 285),     # 时间类达到满分
    ('less', 285, 245, 285),     # 时间类恰好及格
    ('less', 265, 245, 285),     # 时间类区间插值
    ('less', 350, 245, 285),     # 时间类慢于及格
])
def test_scoring_formula_matches_android(direction, value, full, pass_):
    fn = _kt_score_more if direction == 'more' else _kt_score_less
    reference = max(0.0, min(100.0, fn(value, full, pass_)))
    pc = SC._score_more if direction == 'more' else SC._score_less
    assert abs(pc(value, full, pass_) - reference) < 1e-9


@pytest.mark.parametrize('std_name,gender,raw,expected_score', [
    ('1000米跑', '男', '4:05', 100.0),
    ('1000米跑', '男', '4:45', 60.0),
    ('足球', '男', '44.5', 100.0),
    ('足球', '男', '54.1', 60.0),
])
def test_calc_score_end_to_end_zhongkao(std_name, gender, raw, expected_score):
    std = S.find_std(S.ZHONGKAO_2026, std_name)
    assert std is not None, f'中考标准缺少 {std_name}'
    result = SC.calc_score(std, gender, raw)
    assert result['ok'] is True
    assert abs(result['score'] - expected_score) < 0.1


@pytest.mark.parametrize('grade,std_name,raw,expected_score', [
    (1, '肺活量', '1700', 100.0),
    (1, '肺活量', '700', 60.0),
    (5, '50米×8往返跑', "1'36\"", 100.0),
    (5, '50米×8往返跑', "2'00\"", 60.0),
    (7, '1000米跑(男)/800米跑(女)', '4:25', 100.0),   # 初中 7 年级男生 265 秒
])
def test_calc_score_end_to_end_by_grade(grade, std_name, raw, expected_score):
    std = S.find_std(S.get_standards_by_grade(grade), std_name)
    assert std is not None, f'{grade} 年级标准缺少 {std_name}'
    result = SC.calc_score(std, '男', raw)
    assert result['ok'] is True
    assert abs(result['score'] - expected_score) < 0.1


# ---------------------------------------------------------------------------
# 10. 负数成绩拦截（≠ 静默满分）
# ---------------------------------------------------------------------------

@pytest.mark.parametrize('unit,raw', [
    ('次', '-5'),
    ('秒', '-5'),
    ('分秒', '-5'),          # Android 侧此分支尚未拦截（见模块 docstring）
    ('分秒', "-1'00"),
])
def test_negative_value_is_rejected(unit, raw):
    with pytest.raises(ValueError):
        SC.parse_value(raw, unit)


def test_negative_score_does_not_earn_full_mark():
    """回归锚点：负数输入绝不能因 LESS 方向命中 `val <= full` 而拿满分。"""
    std = S.find_std(S.ZHONGKAO_2026, '1000米跑')   # direction=LESS
    for raw in ('-5', "-1'00"):
        result = SC.calc_score(std, '男', raw)
        assert result['ok'] is False, f'负数 {raw} 未被拒绝'
        assert result['score'] is None


def test_invalid_text_returns_error_not_exception():
    std = S.find_std(S.ZHONGKAO_2026, '1000米跑')
    result = SC.calc_score(std, '男', 'abc')
    assert result['ok'] is False and result['score'] is None


# ---------------------------------------------------------------------------
# 11. Android 备份导入时的学段还原（此前硬编码 primary/None）
# ---------------------------------------------------------------------------

@pytest.mark.parametrize('code,ttype,grade,tag', [
    ('1', 'primary', 1, '1年级'),
    ('6', 'primary', 6, '6年级'),
    ('7', 'junior', 7, '初一'),
    ('9', 'junior', 9, '初三'),
    ('10', 'senior', 10, '高一'),
    ('12', 'senior', 12, '高三'),
    ('13', 'zhongkao', None, '中考2026'),
    (7, 'junior', 7, '初一'),        # 数字型也要能处理
    (None, 'primary', None, 'Android导入'),
    ('', 'primary', None, 'Android导入'),
])
def test_archive_type_for_grade(code, ttype, grade, tag):
    assert S.archive_type_for_grade(code) == (ttype, grade, tag)


def test_imported_grade_gets_correct_standard_set():
    """导入还原出的学段必须能取到对应标准（且不是小学标准）。"""
    ttype, grade, _ = S.archive_type_for_grade('10')
    assert ttype == 'senior'
    stds = S.get_standards_by_grade(grade)
    assert stds == S.get_standards_by_grade('10')
    assert stds is not S.PRIMARY.get(10)


# ---------------------------------------------------------------------------
# 12. 学段档案 Excel 生成（此前 PC 只有 primary / zhongkao 两条生成路径）
# ---------------------------------------------------------------------------

@pytest.mark.parametrize('code,name,expect_tag', [
    ('7', '初中生', '初一'),
    ('10', '高中生', '高一'),
])
def test_excel_sheet_generation_for_junior_senior(tmp_path, code, name, expect_tag):
    """验证 _build_grade_sheet 通用化后：学段文案正确、项目齐全。"""
    from excel_builder import append_record
    from openpyxl import load_workbook

    ttype, grade, tag = S.archive_type_for_grade(code)
    _, file_path = append_record({
        'name': name, 'age': 13, 'gender': '男', 'school': '某中学', 'phone': '',
        'date': '2026-09-12', 'table_type': ttype, 'grade': grade, 'sheet_tag': tag,
        'zk_plan': '', 'evaluation': '', 'records': {},
    }, str(tmp_path))

    wb = load_workbook(file_path)
    targets = [n for n in wb.sheetnames if expect_tag in n]
    assert targets, f'未找到含"{expect_tag}"的工作表：{wb.sheetnames}'
    ws = wb[targets[0]]

    assert expect_tag in str(ws['A1'].value), f'标题未体现学段：{ws["A1"].value}'
    col_a = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    std_names = [s.name for s in S.get_standards_by_grade(grade)]
    missing = [n for n in std_names if n not in col_a]
    assert not missing, f'{expect_tag}档案缺少项目：{missing}'
