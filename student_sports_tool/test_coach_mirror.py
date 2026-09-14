# -*- coding: utf-8 -*-
"""v67 测试：PC 教练页「手机备份只读镜像」（李哥拍板 D2）。

背景：教练数据重心在 Android（薪资结算 / 排课 / 团队管理都在手机），
PC 只有 6 个基础字段且不参与同步 → 同一份数据两端各存一份早晚对不上。
改为：PC 从备份 `coaches[]` 解析后只读展示，整体替换、以手机为真源。

本测试锁四条底线：
1. **字段映射** —— Android 只导出 name/phone/specialty/status 四字段，
   其余列留空（不动同步协议）；
2. **状态映射** —— status=='在职' → is_active=True，休假/离职 → False；
3. **幂等** —— 同一份备份重复导入不产生重复行；
4. **以手机为真源** —— 手机端删掉的教练不能在 PC 残留（整体替换）。

另有两条解析底线：meta 无 coaches[]（旧备份）不报错；导出结构可被解析层识别。
"""
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.abspath('.'))
sys.path.insert(0, os.path.join(os.path.abspath('.'), 'data_center'))

import coach_manager as cm                                      # noqa: E402
from data_center.android_backup_parser import (                  # noqa: E402
    _normalize_coach, parse_meta_json,
)


def _mk_dir():
    return tempfile.mkdtemp(prefix='_v67_coach_')


def _mirror(name='张三', phone='13800000000', specialty='田径', status='在职'):
    """模拟 Android BackupManager 导出的 coaches[] 一行（只有 4 个字段）。"""
    return {'name': name, 'phone': phone, 'specialty': specialty, 'status': status}


# === 解析层 ===

def test_normalize_coach_maps_four_exported_fields():
    item = _normalize_coach({'name': ' 李四 ', 'phone': '13900000000',
                             'specialty': '篮球', 'status': '在职'})
    assert item == {'name': '李四', 'phone': '13900000000',
                    'specialty': '篮球', 'status': '在职'}


def test_normalize_coach_status_defaults_to_active():
    """status 缺失时按在职处理（不因缺字段把人显示成离职）。"""
    assert _normalize_coach({'name': '王五'})['status'] == '在职'


def test_parse_meta_json_exposes_coaches():
    p = os.path.join(_mk_dir(), 'export_meta.json')
    with open(p, 'w', encoding='utf-8') as f:
        json.dump({
            'students': [{'name': '张三'}],
            'coaches': [
                {'name': '赵教练', 'phone': '13700000000',
                 'specialty': '体能', 'status': '休假'},
            ],
        }, f, ensure_ascii=False)
    parsed = parse_meta_json(p)
    assert parsed['coaches'][0]['name'] == '赵教练'
    assert parsed['coaches'][0]['status'] == '休假'


def test_parse_meta_json_without_coaches_key_is_empty():
    """旧备份（无 coaches[]）解析不报错，返回空列表。"""
    p = os.path.join(_mk_dir(), 'export_meta.json')
    with open(p, 'w', encoding='utf-8') as f:
        json.dump({'students': [{'name': '张三'}]}, f, ensure_ascii=False)
    assert parse_meta_json(p)['coaches'] == []


# === 镜像落盘 ===

def test_replace_mirror_writes_rows_and_status():
    d = _mk_dir()
    n = cm.replace_mirror(d, [_mirror(), _mirror('李四', status='离职')])
    assert n == 2
    rows = cm.list_coaches(d, include_inactive=True)
    assert [r['name'] for r in rows] == ['张三', '李四']
    assert rows[0]['phone'] == '13800000000'
    assert rows[0]['specialty'] == '田径'
    # 备份未导出的三列留空（不动协议）
    assert rows[0]['role'] == ''
    assert rows[0]['join_date'] == ''
    assert rows[0]['note'] == ''
    # status → is_active
    assert rows[0]['is_active'] is True
    assert rows[1]['is_active'] is False


def test_replace_mirror_is_idempotent():
    d = _mk_dir()
    batch = [_mirror(), _mirror('李四')]
    cm.replace_mirror(d, batch)
    cm.replace_mirror(d, batch)
    assert len(cm.list_coaches(d, include_inactive=True)) == 2


def test_replace_mirror_drops_rows_deleted_on_phone():
    """手机删掉的教练不能在 PC 残留 —— 整体替换语义（这是 D2 的核心目的）。"""
    d = _mk_dir()
    cm.replace_mirror(d, [_mirror(), _mirror('李四')])
    cm.replace_mirror(d, [_mirror()])
    assert [r['name'] for r in cm.list_coaches(d, include_inactive=True)] == ['张三']


def test_replace_mirror_skips_blank_and_invalid_rows():
    """姓名为空/非 dict 的行跳过，不写脏数据。"""
    d = _mk_dir()
    n = cm.replace_mirror(d, [_mirror(), {'name': '   '}, 'not-a-dict', None])
    assert n == 1
    assert [r['name'] for r in cm.list_coaches(d, include_inactive=True)] == ['张三']


def test_replace_mirror_keeps_salary_columns_untouched():
    """回归：镜像替换只写 6 列基础字段，不影响表头结构（PC 不引入薪资列）。"""
    d = _mk_dir()
    cm.replace_mirror(d, [_mirror()])
    from openpyxl import load_workbook
    wb = load_workbook(cm.ensure_file(d))
    ws = wb[cm.COACH_SHEET]
    header = [ws.cell(row=1, column=i).value for i in range(1, len(cm.HEADERS) + 1)]
    assert header == cm.HEADERS
    assert ws.max_column == len(cm.HEADERS)
