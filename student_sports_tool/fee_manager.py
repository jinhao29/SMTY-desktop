# -*- coding: utf-8 -*-
"""管理层：收费记录数据读写与财务自动计算。

职责：
- 维护「收费记录.xlsx」（位于档案目录下，与课时记录.xlsx 并列）
- 单表「明细」：每次收费一条（学员、日期、金额、对应课时数、方式、备注）
- 财务口径（全部自动推导，无手工状态）：
    实收       = Σ 收费金额
    已购课时   = Σ 收费课时数
    课时单价   = 实收 / 已购课时（加权平均，保留 2 位）
    应收       = 学员总课时 × 课时单价（总课时来自课时记录汇总表）
    已消课费用 = 已上课时 × 课时单价（每录一节课自动增加）
    待收       = 应收 - 实收

数据结构：
- 明细表列：序号 | 日期 | 学员 | 金额 | 课时数 | 收款方式 | 备注
"""
import os
from datetime import datetime, date as _date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from file_lock import file_lock, atomic_save_workbook


FEE_FILE = '收费记录.xlsx'

DETAIL_HEADERS = ['序号', '日期', '学员', '金额', '课时数', '收款方式', '备注']

THIN = Side(style='thin', color='888888')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
HEADER_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF', name='微软雅黑')


def _fee_file_path(dir_path):
    return os.path.join(dir_path, FEE_FILE)


def _save_wb(fpath, wb):
    with file_lock(fpath):
        atomic_save_workbook(wb, fpath)


def _load_wb(fpath):
    with file_lock(fpath, mode='r'):
        return load_workbook(fpath)


def _invalidate_meta_index(dir_path: str = ''):
    try:
        from data_center.meta_index_store import invalidate_index
        invalidate_index(dir_path)
    except Exception:
        pass


def _norm_date(v):
    """归一化日期值为 'YYYY-MM-DD' 字符串（明细可能混存 date/datetime/字符串）。"""
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, _date):
        return v.strftime('%Y-%m-%d')
    return str(v or '').strip()[:10]


def _ensure_file(dir_path):
    """确保收费记录文件存在，不存在则创建带表头的空文件。"""
    fpath = _fee_file_path(dir_path)
    if os.path.exists(fpath):
        return fpath
    wb = Workbook()
    ws = wb.active
    ws.title = '明细'
    for i, h in enumerate(DETAIL_HEADERS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER
    for col, width in zip('ABCDEFG', (6, 12, 12, 10, 8, 10, 18)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    _save_wb(fpath, wb)
    return fpath


def _read_detail(wb):
    """读取明细，返回 [{seq,date,name,amount,hours,method,note,row}]。"""
    ws = wb['明细']
    records = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=3).value
        if not name:
            continue
        records.append({
            'seq': ws.cell(row=r, column=1).value,
            'date': _norm_date(ws.cell(row=r, column=2).value),
            'name': str(name).strip(),
            'amount': float(ws.cell(row=r, column=4).value or 0),
            'hours': float(ws.cell(row=r, column=5).value or 0),
            'method': str(ws.cell(row=r, column=6).value or '').strip(),
            'note': str(ws.cell(row=r, column=7).value or '').strip(),
            'row': r,
        })
    return records


def get_payments(dir_path, name=None):
    """读取收费记录。name 指定时只返回该学员的记录（按日期倒序）。"""
    fpath = _ensure_file(dir_path)
    wb = _load_wb(fpath)
    records = _read_detail(wb)
    if name:
        records = [r for r in records if r['name'] == name]
    records.sort(key=lambda x: x['date'], reverse=True)
    return records


def add_payment(dir_path, name, date, amount, hours, method='', note=''):
    """新增一条收费记录。amount/hours 必须 > 0。"""
    if not name or amount <= 0 or hours <= 0:
        return False
    fpath = _ensure_file(dir_path)
    wb = _load_wb(fpath)
    ws = wb['明细']
    target_row = ws.max_row + 1
    if ws.cell(row=ws.max_row, column=1).value is None:
        target_row = ws.max_row
    date_str = _norm_date(date) or datetime.now().strftime('%Y-%m-%d')
    values = [target_row - 1, date_str, name, float(amount), float(hours), method, note]
    for c, v in enumerate(values, 1):
        ws.cell(row=target_row, column=c, value=v)
    _save_wb(fpath, wb)
    _invalidate_meta_index(dir_path)
    try:
        from data_center.sync_beacon import notify_data_changed
        notify_data_changed()
    except Exception:
        pass
    return True


def delete_payment(dir_path, row_num):
    """按工作表行号删除收费记录，并重排序号。"""
    fpath = _ensure_file(dir_path)
    wb = _load_wb(fpath)
    ws = wb['明细']
    if row_num < 2 or row_num > ws.max_row:
        return False
    ws.delete_rows(row_num)
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1, value=r - 1)
    _save_wb(fpath, wb)
    _invalidate_meta_index(dir_path)
    try:
        from data_center.sync_beacon import notify_data_changed
        notify_data_changed()
    except Exception:
        pass
    return True


#==== 财务计算（纯函数，便于测试） ====

def unit_price(paid_amount, paid_hours):
    """加权平均课时单价：实收 / 已购课时，保留 2 位；无数据返回 0.0。"""
    if paid_hours <= 0:
        return 0.0
    return round(paid_amount / paid_hours, 2)


def compute_student_finance(total, attended, paid_amount, paid_hours):
    """单学员财务指标。返回 dict（含单价/应收/已消课费用/待收/状态）。"""
    price = unit_price(paid_amount, paid_hours)
    receivable = round(total * price, 2)
    consumed = round(attended * price, 2)
    due = round(receivable - paid_amount, 2)
    if paid_amount <= 0:
        status = '未收费'
    elif due <= 0.005:
        status = '已结清'
    else:
        status = '待收款'
    return {
        'unit_price': price,
        'receivable': receivable,
        'consumed': consumed,
        'paid': round(paid_amount, 2),
        'due': due,
        'status': status,
    }


def compute_finance(dir_path, lesson_summaries=None, fee_records=None):
    """汇总所有学员财务数据。

    lesson_summaries: lesson_manager.get_summary() 结果（None 则内部读取）
    fee_records:      get_payments() 结果（None 则内部读取）
    返回按学员名排序的列表，每条：
    {name,total,attended,remaining, unit_price,receivable,consumed,paid,due,status}
    """
    if lesson_summaries is None:
        from lesson_manager import get_summary
        lesson_summaries = get_summary(dir_path)
    if fee_records is None:
        fee_records = get_payments(dir_path)

    pay_map = {}
    for rec in fee_records:
        agg = pay_map.setdefault(rec['name'], {'amount': 0.0, 'hours': 0.0})
        agg['amount'] += rec['amount']
        agg['hours'] += rec['hours']

    result = []
    for s in lesson_summaries:
        agg = pay_map.pop(s['name'], {'amount': 0.0, 'hours': 0.0})
        row = dict(s)
        row.update(compute_student_finance(
            s.get('total', 0), s.get('attended', 0),
            agg['amount'], agg['hours']))
        result.append(row)
    # 有收费但课时汇总里已无此学员的（如已删档），仍列出收费情况
    for name, agg in pay_map.items():
        row = {'name': name, 'total': 0, 'attended': 0, 'remaining': -agg['hours'],
               'last_date': '', 'note': ''}
        row.update(compute_student_finance(0, 0, agg['amount'], agg['hours']))
        row['status'] = '已结清' if agg['amount'] > 0 else '未收费'
        result.append(row)
    result.sort(key=lambda x: x['name'])
    return result


def compute_totals(finance_rows, fee_records=None, today=None):
    """顶部统计卡数据。返回 {total_paid,total_receivable,total_due,month_paid}。

    month_paid 需要收费明细（fee_records，get_payments() 结果）过滤当月；
    传入 None 时当月实收计 0。
    """
    today = today or _date.today()
    month_prefix = today.strftime('%Y-%m')
    month_paid = 0.0
    if fee_records:
        month_paid = sum(r['amount'] for r in fee_records
                         if str(r['date']).startswith(month_prefix))
    return {
        'total_paid': round(sum(r['paid'] for r in finance_rows), 2),
        'total_receivable': round(sum(r['receivable'] for r in finance_rows), 2),
        'total_due': round(sum(r['due'] for r in finance_rows), 2),
        'month_paid': round(month_paid, 2),
        'month_prefix': month_prefix,
    }
