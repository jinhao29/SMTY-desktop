# -*- coding: utf-8 -*-
"""协调层：导入/导出/备份/恢复流程入口。

职责：
- 接收 UI 层的操作请求，编排备份/恢复/导入/导出流程
- 备份创建 / 恢复 / 列表管理 / 清理已下沉到 backup/ 子包（P2 超大文件拆分），
  本模块仅保留导入/导出，并 re-export 备份接口，对外接口保持不变。
"""
import os
import json

from openpyxl import Workbook, load_workbook
from file_lock import file_lock, atomic_save_workbook
from data_exporter import export_summary_excel

# 备份接口 re-export（对外接口不变）
from backup.backup_creator import do_backup, do_auto_backup
from backup.backup_restorer import do_restore

# 小程序桥（阶段五互通）：入口统一走 backup_coordinator
from data_center.miniprogram_bridge import (
    import_miniprogram_backup, export_miniprogram_backup,
)

__all__ = [
    'do_backup', 'do_auto_backup', 'do_restore',
    'do_export_summary', 'import_students_from_excel', 'generate_import_template',
    'import_miniprogram_backup', 'export_miniprogram_backup',
]


def do_export_summary(dir_path, output_path, progress_cb=None):
    """导出学员汇总数据为 Excel。"""
    dir_path = os.path.normpath(dir_path)
    output_path = os.path.normpath(output_path)
    if progress_cb:
        progress_cb('正在汇总学员数据...')
    count = export_summary_excel(dir_path, output_path)
    if progress_cb:
        progress_cb(f'导出完成：{count} 位学员')
    return count


def import_students_from_excel(import_path, dir_path, progress_cb=None):
    """从 Excel 批量导入学员名单，创建空档案。

    参数:
        import_path: 导入的 Excel 文件路径
        dir_path: 档案目录

    返回: (成功数, 冲突列表)
    """
    import_path = os.path.normpath(import_path)
    dir_path = os.path.normpath(dir_path)
    if not os.path.exists(import_path):
        raise FileNotFoundError(f'导入文件不存在：{import_path}')
    wb = load_workbook(import_path, read_only=True)
    ws = wb.active
    created = 0
    conflicts = []

    # 解析表头定位列（支持中英文表头）
    headers = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=c).value
        if val:
            headers[str(val).strip()] = c

    name_col = headers.get('姓名') or headers.get('name') or 1
    gender_col = headers.get('性别') or headers.get('gender')
    school_col = headers.get('学校') or headers.get('school')
    phone_col = headers.get('联系电话') or headers.get('phone') or headers.get('电话')

    os.makedirs(dir_path, exist_ok=True)
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=name_col).value
        if not name:
            continue
        name = str(name).strip()
        file_path = os.path.join(dir_path, f'{name}.xlsx')
        if os.path.exists(file_path):
            conflicts.append(name)
            if progress_cb:
                progress_cb(f'跳过已存在：{name}')
            continue
        # 创建空档案
        wb_new = Workbook()
        ws_meta = wb_new.active
        ws_meta.title = '_meta'
        meta = {
            'info': {
                'name': name,
                'gender': str(ws.cell(row=r, column=gender_col).value or '男').strip() if gender_col else '男',
                'school': str(ws.cell(row=r, column=school_col).value or '').strip() if school_col else '',
                'phone': str(ws.cell(row=r, column=phone_col).value or '').strip() if phone_col else '',
            },
            'records': [],
        }
        ws_meta.cell(row=1, column=1, value=json.dumps(meta, ensure_ascii=False))
        ws_meta.sheet_state = 'hidden'
        wb_new.save(file_path)
        created += 1
        if progress_cb:
            progress_cb(f'已创建档案：{name}')
    wb.close()
    if progress_cb:
        msg = f'导入完成：新建 {created} 个档案'
        if conflicts:
            msg += f'，{len(conflicts)} 个已存在被跳过'
        progress_cb(msg)
    return created, conflicts


def generate_import_template(output_path):
    """生成导入用 Excel 模板。"""
    output_path = os.path.normpath(output_path)
    wb = Workbook()
    ws = wb.active
    ws.title = '学员名单'
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    thin = Side(style='thin', color='888888')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color='FFFFFF', name='微软雅黑')
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')

    headers = ['姓名', '性别', '学校', '联系电话']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border
    # 示例行
    samples = [('张三', '男', '深圳实验学校', '13800000001'),
               ('李四', '女', '深圳中学', '13800000002')]
    for r, row in enumerate(samples, start=2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = border
            cell.alignment = center
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 16
    with file_lock(output_path):
        atomic_save_workbook(wb, output_path)
