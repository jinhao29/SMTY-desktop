# -*- coding: utf-8 -*-
"""管理层：进步对比汇总表生成。

职责：
- 从 _meta 元数据重建「进步对比」工作表（置顶）
- 矩阵布局：行=测试项目，列=历次测评
- 单元格显示「成绩(得分)」，末列显示进步幅度（本次 vs 首次）
- 底部列出总分与历次评价
- 仅汇总成绩型档案（primary/zhongkao/zhongkao_old）
"""
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill

THIN = Side(style='thin', color='888888')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
TITLE_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
HEADER_FILL = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
SCORE_FILL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
PROGRESS_FILL = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
TITLE_FONT = Font(bold=True, size=12, color='FFFFFF', name='微软雅黑')
HEADER_FONT = Font(bold=True, name='微软雅黑')
NAME_FONT = Font(bold=True, size=14, name='微软雅黑')
SUMMARY_NAME = '进步对比'


def _safe_merge(ws, ranges):
    """安全合并单元格，忽略合并冲突。"""
    for m in ranges:
        try:
            ws.merge_cells(start_row=m[0], start_column=m[1], end_row=m[2], end_column=m[3])
        except Exception:
            pass


def _style_block(ws, r1, c1, r2, c2, fill=None, font=None, align=CENTER):
    """为矩形区域统一设置边框、对齐、填充、字体。"""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = align
            if fill:
                cell.fill = fill
            if font:
                cell.font = font


def _progress_text(diff):
    """数值差转进步文本（↑正/↓负/→持平）。"""
    if diff is None:
        return ''
    if diff > 0.05:
        return f"↑{diff:.1f}"
    if diff < -0.05:
        return f"↓{abs(diff):.1f}"
    return '→0.0'


def _collect_items(records):
    """从历次记录收集项目并集（保持首次出现顺序）。"""
    items = []
    for rec in records:
        for name in rec.get('scores', {}).keys():
            if name not in items:
                items.append(name)
    return items


def _cell_text(score_rec):
    """单条成绩记录转单元格文本：成绩(得分分)。"""
    if not score_rec:
        return ''
    val = score_rec.get('value', '')
    score = score_rec.get('score')
    if score is None:
        return str(val)
    if isinstance(score, float) and score == int(score):
        score_txt = f"{int(score)}"
    else:
        score_txt = f"{score:.1f}"
    return f"{val}\n({score_txt}分)" if val else f"{score_txt}分"


def build_summary_sheet(wb, meta):
    """根据元数据重建「进步对比」汇总表（置顶）。

    参数:
        wb: openpyxl Workbook
        meta: {'info': {...}, 'records': [{seq,date,tag,scores,total,evaluation}, ...]}
    """
    if SUMMARY_NAME in wb.sheetnames:
        del wb[SUMMARY_NAME]
    ws = wb.create_sheet(title=SUMMARY_NAME, index=0)

    info = meta.get('info', {})
    records = meta.get('records', [])
    n = len(records)

    # 列宽
    ws.column_dimensions['A'].width = 18
    for i in range(n + 1):
        col_letter = chr(ord('B') + i)
        ws.column_dimensions[col_letter].width = 14

    last_col = 2 + n  # 项目列 + N次测评 + 进步列
    r = 1

    # 主标题
    ws.cell(row=r, column=1, value=f"{info.get('name', '')}  进步对比汇总")
    _safe_merge(ws, [(r, 1, r, last_col)])
    _style_block(ws, r, 1, r, last_col, fill=TITLE_FILL, font=NAME_FONT, align=LEFT)
    r += 1

    # 基本信息行
    info_txt = (f"性别：{info.get('gender', '—')}    "
                f"学校：{info.get('school', '—')}    "
                f"电话：{info.get('phone', '—')}    "
                f"累计测评：{n} 次")
    ws.cell(row=r, column=1, value=info_txt)
    _safe_merge(ws, [(r, 1, r, last_col)])
    _style_block(ws, r, 1, r, last_col, align=LEFT)
    r += 2

    if n == 0:
        ws.cell(row=r, column=1, value='（暂无测评记录）')
        _safe_merge(ws, [(r, 1, r, last_col)])
        _style_block(ws, r, 1, r, last_col, align=LEFT)
        return

    # 表头三行：序号 / 日期 / 类型
    ws.cell(row=r, column=1, value='项目')
    ws.cell(row=r + 1, column=1, value='日期')
    ws.cell(row=r + 2, column=1, value='类型')
    for i, rec in enumerate(records):
        col = 2 + i
        ws.cell(row=r, column=col, value=f"第{rec.get('seq', i + 1)}次")
        ws.cell(row=r + 1, column=col, value=rec.get('date', ''))
        ws.cell(row=r + 2, column=col, value=rec.get('tag', ''))
    progress_col = 2 + n
    ws.cell(row=r, column=progress_col, value='进步')
    ws.cell(row=r + 1, column=progress_col, value='本次vs首次')
    ws.cell(row=r + 2, column=progress_col, value='↑提升 /↓下降')
    _style_block(ws, r, 1, r + 2, progress_col, fill=HEADER_FILL, font=HEADER_FONT)
    r += 3

    # 项目数据行
    items = _collect_items(records)
    for item_name in items:
        ws.cell(row=r, column=1, value=item_name).font = HEADER_FONT
        first_score = None
        last_score = None
        for i, rec in enumerate(records):
            col = 2 + i
            sc = rec.get('scores', {}).get(item_name)
            ws.cell(row=r, column=col, value=_cell_text(sc))
            if sc and sc.get('score') is not None:
                if first_score is None:
                    first_score = sc['score']
                last_score = sc['score']
        # 进步列（得分差）
        if first_score is not None and last_score is not None:
            diff = last_score - first_score
            cell = ws.cell(row=r, column=progress_col, value=_progress_text(diff))
            cell.fill = PROGRESS_FILL
            cell.font = Font(bold=True, color='2E7D32' if diff >= 0 else 'C62828', name='微软雅黑')
        _style_block(ws, r, 1, r, progress_col)
        ws.row_dimensions[r].height = 32
        r += 1

    # 总分行
    ws.cell(row=r, column=1, value='总分').font = HEADER_FONT
    first_total = None
    last_total = None
    for i, rec in enumerate(records):
        col = 2 + i
        total = rec.get('total')
        if total is not None:
            ws.cell(row=r, column=col, value=round(total, 1))
            if first_total is None:
                first_total = total
            last_total = total
    if first_total is not None and last_total is not None:
        diff = last_total - first_total
        cell = ws.cell(row=r, column=progress_col, value=_progress_text(diff))
        cell.font = Font(bold=True, color='2E7D32' if diff >= 0 else 'C62828', name='微软雅黑')
    _style_block(ws, r, 1, r, progress_col, fill=SCORE_FILL,
                 font=Font(bold=True, name='微软雅黑'))
    ws.row_dimensions[r].height = 26
    r += 2

    # 历次评价区
    ws.cell(row=r, column=1, value='历次评价与建议').font = TITLE_FONT
    _safe_merge(ws, [(r, 1, r, progress_col)])
    _style_block(ws, r, 1, r, progress_col, fill=TITLE_FILL, align=LEFT)
    r += 1
    for rec in records:
        label = f"第{rec.get('seq', '?')}次 {rec.get('date', '')} ({rec.get('tag', '')})"
        ws.cell(row=r, column=1, value=label).font = HEADER_FONT
        ws.cell(row=r, column=2, value=rec.get('evaluation', '') or '（无）')
        _safe_merge(ws, [(r, 2, r, progress_col)])
        _style_block(ws, r, 1, r, progress_col, align=LEFT)
        ws.row_dimensions[r].height = 28
        r += 1

    ws.freeze_panes = 'B5'
