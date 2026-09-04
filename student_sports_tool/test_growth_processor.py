# -*- coding: utf-8 -*-
"""growth_processor 与 profile_storage 扩展字段的自检测试。

运行：pytest test_growth_processor.py -v
（无 pytest 时也可直接 python test_growth_processor.py）
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'student_profile'))

import growth_processor as gp
import profile_storage as ps
import profile_manager as pm


# ==================== 身高预测 ====================

def test_cmh_formula():
    # 男孩 = (父+母+13)/2；女孩 = (父+母-13)/2
    r = gp.predict_height('男', 10, 175, 160, 0, 0, 0)
    assert r['target_height'] == 174.0, r
    r = gp.predict_height('女', 10, 175, 160, 0, 0, 0)
    assert r['target_height'] == 161.0, r


def test_missing_parents_returns_none():
    assert gp.predict_height('男', 10, 0, 160, 8, 4, 200) is None
    assert gp.predict_height('男', 10, None, None, 0, 0, 0) is None


def test_adjustments_and_clamp():
    # 睡眠9h(+1) 营养5(+1) 运动200min(+1.5) → +3.5
    r = gp.predict_height('男', 10, 175, 160, 9.0, 5, 200)
    assert r['adjustment'] == 3.5
    assert r['adjusted_height'] == 177.5
    # 超限钳制：睡眠<7(-1) 营养<=2(-1) 运动不足(0) → -2
    r = gp.predict_height('男', 10, 175, 160, 6.0, 2, 60)
    assert r['adjustment'] == -2.0
    # 未填写不修正
    r = gp.predict_height('男', 10, 175, 160, 0, 0, 0)
    assert r['adjustment'] == 0.0
    assert '保持当前生活习惯' in r['advice_text']


def test_advice_text():
    r = gp.predict_height('男', 10, 175, 160, 6.0, 2, 60)
    assert '睡眠' in r['advice_text'] and '营养' in r['advice_text'] and '运动' in r['advice_text']
    r = gp.predict_height('男', 10, 175, 160, 9.0, 5, 300)
    assert '后天环境优秀' in r['advice_text']


def test_bone_age_warning():
    assert gp.predict_height('女', 15, 170, 158, 0, 0, 0)['bone_age_warning']
    assert gp.predict_height('男', 17, 175, 160, 0, 0, 0)['bone_age_warning']
    assert gp.predict_height('男', 16, 175, 160, 0, 0, 0)['bone_age_warning'] is None
    assert gp.predict_height('女', 10, 170, 158, 0, 0, 0)['bone_age_warning'] is None


def test_rating_attached():
    # 男孩 10 岁预测 177.5 > P97 151.7 → 优秀
    r = gp.predict_height('男', 10, 175, 160, 9.0, 5, 200)
    assert r['rating']['rating'] == 'tall'
    # 年龄超范围 → 不评级
    r = gp.predict_height('男', 20, 175, 160, 0, 0, 0)
    assert r['rating'] is None


# ==================== 身高评级 ====================

def test_rate_height():
    assert gp.rate_height('男', 10, 128.0)['rating'] == 'short'   # < P3 128.7
    assert gp.rate_height('男', 10, 140.2)['rating'] == 'average'  # P3~P97
    assert gp.rate_height('女', 8, 139.0)['rating'] == 'tall'      # > P97 138.2
    assert gp.rate_height('男', 2, 90.0) is None                   # 年龄超范围
    assert gp.rate_height('男', 10, 0) is None                     # 身高无效


# ==================== TDEE ====================

def test_tdee_mifflin():
    # 男 80kg/180cm/30岁 久坐1.2：BMR=1780 TDEE=2136 成人缺口400
    r = gp.calc_tdee('男', 80, 180, 30, 'sedentary')
    assert r['bmr'] == 1780 and r['tdee'] == 2136
    assert r['deficit_advice'] == 400 and r['warning_text'] is None
    # 女 40kg/150cm/12岁 中度1.55：BMR=1116.5→1117（对齐 Kotlin roundToInt）
    r = gp.calc_tdee('女', 40, 150, 12, 'moderate')
    assert r['bmr'] == 1117 and r['tdee'] == 1731
    assert r['deficit_advice'] is None and '热量缺口' in r['warning_text']


def test_tdee_invalid():
    assert gp.calc_tdee('男', 0, 180, 30) is None
    assert gp.calc_tdee('男', 80, 0, 30) is None
    assert gp.calc_tdee('男', 80, 180, 0) is None


# ==================== 膳食模板 ====================

def test_diet_templates():
    assert len(gp.DIET_TEMPLATES) == 3
    ids = {t['id'] for t in gp.DIET_TEMPLATES}
    assert ids == {'tpl_regular', 'tpl_training', 'tpl_fat_loss'}
    for t in gp.DIET_TEMPLATES:
        # 3+2：三餐 + 两次加餐，5 个餐次齐全
        assert set(t['meals'].keys()) == {k for k, _ in gp.MEAL_ORDER}
        assert t['pre_workout_tip'] and t['post_workout_tip']
    assert gp.get_template('tpl_regular')['name'] == '常规健康发育型'
    assert gp.get_template('nope') is None


# ==================== profile_storage 扩展字段 + 体型历史 ====================

def _tmp_dir():
    return tempfile.mkdtemp(prefix='smty_test_')


def test_extended_fields_roundtrip():
    d = _tmp_dir()
    p = pm.save_student(d, '张三', 12, 150.0, 40.0, '小学六年级', note='测试',
                        gender='女', school='实验小学', phone='13800000000',
                        father_height=175, mother_height=160,
                        sleep_hours=8.5, nutrition_score=4, sports_mins=200)
    assert p['gender'] == '女' and p['school'] == '实验小学'
    assert p['father_height'] == 175.0 and p['mother_height'] == 160.0
    students = pm.list_students(d)
    s = next(x for x in students if x['name'] == '张三')
    assert s['phone'] == '13800000000' and s['sleep_hours'] == 8.5
    assert s['nutrition_score'] == 4 and s['sports_mins'] == 200
    # 空值默认
    pm.save_student(d, '李四', 10, None, None, '小学四年级')
    s = next(x for x in pm.list_students(d) if x['name'] == '李四')
    assert s['gender'] == '男' and s['father_height'] is None


def test_old_file_migration():
    """旧 8 列档案文件：首次写入自动补齐扩展列，原字段不丢失。"""
    d = _tmp_dir()
    # 先用旧表头造一个 8 列文件
    from openpyxl import Workbook
    fpath = os.path.join(d, ps.PROFILE_FILE)
    wb = Workbook()
    ws = wb.active
    ws.title = ps.PROFILE_SHEET
    for i, h in enumerate(ps.PROFILE_HEADERS, 1):
        ws.cell(row=1, column=i, value=h)
    for i, v in enumerate(['王五', 11, 145.0, 35.0, 16.6, '偏瘦', '小学五年级', '旧数据'], 1):
        ws.cell(row=2, column=i, value=v)
    wb.save(fpath)
    # upsert 扩展字段 → 旧数据保留 + 新列补齐（备注不传则按空覆盖，为既有行为）
    pm.save_student(d, '王五', 11, 146.0, 36.0, '小学五年级', note='旧数据', gender='男',
                    school='旧校', father_height=170, mother_height=155)
    rows = ps.read_all(d)
    s = next(x for x in rows if x['name'] == '王五')
    assert s['note'] == '旧数据' and s['grade'] == '小学五年级'
    assert s['gender'] == '男' and s['school'] == '旧校'
    assert s['father_height'] == 170.0 and s['mother_height'] == 155.0
    # 读取旧 8 列文件（未迁移）也应优雅降级
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = ps.PROFILE_SHEET
    for i, h in enumerate(ps.PROFILE_HEADERS, 1):
        ws2.cell(row=1, column=i, value=h)
    for i, v in enumerate(['赵六', 9, 135.0, 30.0, 16.5, '偏瘦', '小学三年级', ''], 1):
        ws2.cell(row=2, column=i, value=v)
    wb2.save(os.path.join(d, ps.PROFILE_FILE))
    s = next(x for x in ps.read_all(d) if x['name'] == '赵六')
    assert s['gender'] == '男' and s['father_height'] is None


def test_body_metric_history():
    d = _tmp_dir()
    pm.save_student(d, '张三', 12, 150.0, 40.0, '小学六年级')
    pm.save_student(d, '张三', 12, 151.0, 41.0, '小学六年级')  # 同日再测 → 覆盖
    today = __import__('datetime').date.today().strftime('%Y-%m-%d')
    history = ps.read_body_metrics(d, '张三')
    assert len(history) == 1 and history[0]['date'] == today
    assert history[0]['height'] == 151.0 and history[0]['weight'] == 41.0
    assert history[0]['bmi'] == round(41.0 / (1.51 ** 2), 1)
    # 不同日期 → 追加
    ps.append_body_metric(d, '张三', '2026-08-01', 148.0, 38.0)
    history = ps.read_body_metrics(d, '张三')
    assert len(history) == 2 and history[0]['date'] == '2026-08-01'
    # 他人数据隔离 + 无效值跳过
    ps.append_body_metric(d, '李四', today, 140.0, 35.0)
    assert len(ps.read_body_metrics(d, '张三')) == 2
    assert ps.append_body_metric(d, '张三', today, 0, 0) is False


def test_student_extra():
    d = _tmp_dir()
    pm.save_student(d, '张三', 12, 150.0, 40.0, '小学六年级')
    assert pm.get_student_extra(d, '张三', 'diet_plan') is None
    assert pm.set_student_extra(d, '张三', 'diet_plan', 'tpl_training')
    assert pm.get_student_extra(d, '张三', 'diet_plan') == 'tpl_training'
    # 软删除状态不受扩展元数据影响
    assert ps.read_all(d)[0]['is_active'] is True


if __name__ == '__main__':
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith('test_') and callable(fn):
            try:
                fn()
                print(f'PASS {name}')
            except AssertionError as e:
                fails += 1
                print(f'FAIL {name}: {e}')
    sys.exit(1 if fails else 0)
