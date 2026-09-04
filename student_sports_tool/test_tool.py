# -*- coding: utf-8 -*-
"""测试评分逻辑与 Excel 生成（pytest 格式，支持 headless / CI 运行）。

运行：python -m pytest test_tool.py -v
"""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from standards import get_primary_standards, get_zhongkao_standards, find_std
from scorer import calc_score


def _score(grade_standards, project, gender, value):
    """按年级标准 + 项目 + 性别 + 成绩值计算得分。"""
    std = find_std(grade_standards, project)
    return calc_score(std, gender, value)


def test_scoring():
    """评分逻辑断言：满分 / 及格 / 各年级 / 中考项目。"""
    # 1年级男生 肺活量 1700(满分) → 100
    assert abs(_score(get_primary_standards(1), '肺活量', '男', '1700')['score'] - 100) < 0.1
    # 1年级男生 肺活量 700(及格) → 60
    assert abs(_score(get_primary_standards(1), '肺活量', '男', '700')['score'] - 60) < 0.1
    # 5年级男生 50米×8 1'36"(满分96秒) → 100
    assert abs(_score(get_primary_standards(5), '50米×8往返跑', '男', "1'36\"")['score'] - 100) < 0.1
    # 5年级男生 50米×8 2'00"(及格120秒) → 60
    assert abs(_score(get_primary_standards(5), '50米×8往返跑', '男', "2'00\"")['score'] - 60) < 0.1
    # 中考 1000米 男 4:05(满分245秒) → 100
    assert abs(_score(get_zhongkao_standards(), '1000米跑', '男', "4:05")['score'] - 100) < 0.1
    # 中考 实心球 男 9.4(满分) → 100
    assert abs(_score(get_zhongkao_standards(), '投掷实心球(2kg)', '男', '9.4')['score'] - 100) < 0.1
    # 中考 足球 男 44.5(满分) → 100
    assert abs(_score(get_zhongkao_standards(), '足球', '男', '44.5')['score'] - 100) < 0.1


def test_home_overview():
    """首页概览统计（compute_overview 纯函数）：分类计数与剩余课时合计。"""
    from home_page import compute_overview
    students = [
        {'name': '甲', 'warn_level': 'red', 'remaining': 2},
        {'name': '乙', 'warn_level': 'yellow', 'remaining': 6},
        {'name': '丙', 'warn_level': 'normal', 'remaining': 10},
        {'name': '丁', 'warn_level': 'unknown', 'remaining': None},
    ]
    assert compute_overview(students) == {
        'total': 4, 'red': 1, 'yellow': 1, 'remaining': 18,
    }
    assert compute_overview([]) == {'total': 0, 'red': 0, 'yellow': 0, 'remaining': 0}


def test_excel_generation(tmp_path):
    """Excel 生成：写入 pytest 临时目录，不再依赖硬编码绝对路径。"""
    from excel_builder import append_record
    from openpyxl import load_workbook

    dir_path = str(tmp_path)

    # 学员1: 小学5年级男生
    s1 = {
        'name': '张三', 'gender': '男', 'school': '深圳实验小学', 'phone': '13800000000',
        'date': '2026-06-29', 'table_type': 'primary', 'grade': 5, 'sheet_tag': '5年级',
        'records': {
            '肺活量': {'value': '2800', 'score': 80.0, 'grade': '良好'},
            '50米跑': {'value': '9.5', 'score': 55.0, 'grade': '不及格'},
            '坐位体前屈': {'value': '16.5', 'score': 100.0, 'grade': '优秀'},
        },
        'evaluation': '柔韧性好，速度需加强，建议加强短跑训练。',
    }
    _, file1 = append_record(s1, dir_path)

    # 学员2: 中考女生
    s2 = {
        'name': '李四', 'gender': '女', 'school': '深圳中学', 'phone': '13900000000',
        'date': '2026-06-29', 'table_type': 'zhongkao', 'grade': None, 'sheet_tag': '中考2026',
        'zk_plan': '2026新方案',
        'records': {
            '1000米跑': {'value': "4'20\"", 'score': 80.0, 'grade': '良好'},
            '投掷实心球(2kg)': {'value': '7.0', 'score': 85.0, 'grade': '良好'},
            '足球': {'value': '45.0', 'score': 98.0, 'grade': '优秀'},
        },
        'evaluation': '整体良好，长跑需加强。',
    }
    _, file2 = append_record(s2, dir_path)

    assert os.path.exists(file1)
    assert os.path.exists(file2)
    wb = load_workbook(file1)
    assert len(wb.sheetnames) >= 1
