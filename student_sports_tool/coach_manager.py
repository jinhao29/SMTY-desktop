# -*- coding: utf-8 -*-
"""数据层：教练档案 Excel 读写。

存储文件：档案目录下的「教练档案.xlsx」
工作表：教练信息
列结构：姓名 | 电话 | 角色 | 专长 | 入职日期 | 备注

设计约定：
- 教练姓名为主键，重名视为同一教练（更新而非新增）
- 软删除：通过隐藏 _meta 工作表的 is_active 字段标记离职状态，
  与学员档案的停用/恢复模式保持一致
- 硬删除（delete_coach）保留给历史上 PC 端手工维护的场景；
  v67 起 PC 教练页改为「手机备份只读镜像」（replace_mirror），不再提供 PC 端新增/编辑入口
"""
import json
import os
import sys
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# 引入统一文件锁（位于当前目录）
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)
from file_lock import file_lock, with_retry, atomic_save_workbook

COACH_FILE = '教练档案.xlsx'
COACH_SHEET = '教练信息'
META_SHEET = '_meta'

HEADERS = ['姓名', '电话', '角色', '专长', '入职日期', '备注']

# 角色建议选项（角色列允许自由填写，下拉仅提供建议）
ROLE_OPTIONS = ['主教练', '助理教练', '体能教练', '兼职教练']

THIN = Side(style='thin', color='888888')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
HEADER_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF', name='微软雅黑')


def _coach_file_path(dir_path: str) -> str:
    """返回教练档案文件完整路径。"""
    return os.path.join(dir_path, COACH_FILE)


def ensure_file(dir_path: str) -> str:
    """确保教练档案文件存在，不存在则创建带表头的空文件。返回文件路径。"""
    fpath = _coach_file_path(dir_path)
    if os.path.exists(fpath):
        return fpath
    os.makedirs(dir_path, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = COACH_SHEET
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER
    widths = [14, 18, 12, 22, 14, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'
    _ensure_meta_sheet(wb)
    atomic_save_workbook(wb, fpath)
    return fpath


def _ensure_meta_sheet(wb) -> None:
    """确保 _meta 隐藏工作表存在：{"coaches": {"姓名": {"is_active": true}}}。"""
    if META_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(title=META_SHEET)
        ws.cell(row=1, column=1, value=json.dumps({'coaches': {}}, ensure_ascii=False))
        ws.sheet_state = 'hidden'


def _read_meta(wb) -> dict:
    """读取 _meta，返回 {'coaches': {...}}；缺失/损坏时返回空结构。"""
    if META_SHEET not in wb.sheetnames:
        return {'coaches': {}}
    raw = wb[META_SHEET].cell(row=1, column=1).value or '{}'
    try:
        meta = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {'coaches': {}}
    meta.setdefault('coaches', {})
    return meta


def _write_meta(wb, meta: dict) -> None:
    """整表重写 _meta（先删后建，避免残留）。"""
    if META_SHEET in wb.sheetnames:
        del wb[META_SHEET]
    ws = wb.create_sheet(title=META_SHEET)
    ws.cell(row=1, column=1, value=json.dumps(meta, ensure_ascii=False))
    ws.sheet_state = 'hidden'


def _now_ms() -> int:
    """当前毫秒时间戳（与学员档案 LWW 时间戳同源）。"""
    return int(time.time() * 1000)


def list_coaches(dir_path: str, include_inactive: bool = False) -> list:
    """读取教练列表。

    参数:
        include_inactive: 是否包含已离职（is_active=False）的教练
    返回:
        [{'name','phone','role','specialty','join_date','note','is_active','row'}]
    """
    fpath = ensure_file(dir_path)

    def _do_read():
        with file_lock(fpath, mode='r', timeout=5.0):
            wb = load_workbook(fpath)
            ws = wb[COACH_SHEET]
            meta = _read_meta(wb)
            coaches_meta = meta.get('coaches', {})
            result = []
            for r in range(2, ws.max_row + 1):
                name = ws.cell(row=r, column=1).value
                if not name or not str(name).strip():
                    continue
                name = str(name).strip()
                coach_meta = coaches_meta.get(name, {})
                is_active = coach_meta.get('is_active', True)
                if not is_active and not include_inactive:
                    continue
                result.append({
                    'name': name,
                    'phone': str(ws.cell(row=r, column=2).value or '').strip(),
                    'role': str(ws.cell(row=r, column=3).value or '').strip(),
                    'specialty': str(ws.cell(row=r, column=4).value or '').strip(),
                    'join_date': str(ws.cell(row=r, column=5).value or '').strip(),
                    'note': str(ws.cell(row=r, column=6).value or '').strip(),
                    'is_active': is_active,
                    'row': r,
                })
            return result

    return with_retry(_do_read, max_retries=3, retry_interval=1.0)


def save_coach(dir_path: str, coach: dict) -> str:
    """新增或更新教练（姓名主键 upsert）。

    coach 必含 'name'；phone/role/specialty/join_date/note 可选。
    返回文件路径；姓名为空抛 ValueError。
    """
    name = str(coach.get('name') or '').strip()
    if not name:
        raise ValueError('教练姓名不能为空')
    fpath = ensure_file(dir_path)
    values = [
        name,
        str(coach.get('phone') or '').strip(),
        str(coach.get('role') or '').strip(),
        str(coach.get('specialty') or '').strip(),
        str(coach.get('join_date') or '').strip(),
        str(coach.get('note') or '').strip(),
    ]

    def _do_save():
        with file_lock(fpath, timeout=5.0):
            wb = load_workbook(fpath)
            ws = wb[COACH_SHEET]
            target_row = None
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(row=r, column=1).value or '').strip() == name:
                    target_row = r
                    break
            if target_row is None:
                target_row = ws.max_row + 1
            for i, v in enumerate(values, 1):
                ws.cell(row=target_row, column=i, value=v)
            meta = _read_meta(wb)
            # 新增教练默认启用（保留既有离职标记：编辑不改变状态）
            meta['coaches'].setdefault(name, {})
            atomic_save_workbook(wb, fpath)
            return True

    with_retry(_do_save, max_retries=3, retry_interval=1.0)
    return fpath


def set_coach_active(dir_path: str, name: str, is_active: bool) -> bool:
    """设置教练在职/离职状态（软删除/恢复）。返回 False 表示未找到。"""
    fpath = ensure_file(dir_path)

    def _do_set():
        with file_lock(fpath, timeout=5.0):
            wb = load_workbook(fpath)
            ws = wb[COACH_SHEET]
            found = False
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(row=r, column=1).value or '').strip() == name:
                    found = True
                    break
            if not found:
                return False
            meta = _read_meta(wb)
            meta.setdefault('coaches', {})
            coach_meta = meta['coaches'].setdefault(name, {})
            coach_meta['is_active'] = bool(is_active)
            coach_meta['updated_at'] = _now_ms()
            _write_meta(wb, meta)
            atomic_save_workbook(wb, fpath)
            return True

    return with_retry(_do_set, max_retries=3, retry_interval=1.0)

def delete_coach(dir_path: str, name: str) -> bool:
    """硬删除教练（v23.9）：移除教练行与 _meta 条目。返回 False 表示未找到。"""
    fpath = ensure_file(dir_path)

    def _do_delete():
        with file_lock(fpath, timeout=5.0):
            wb = load_workbook(fpath)
            ws = wb[COACH_SHEET]
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(row=r, column=1).value or '').strip() == name:
                    ws.delete_rows(r)
                    meta = _read_meta(wb)
                    meta.get('coaches', {}).pop(name, None)
                    _write_meta(wb, meta)
                    atomic_save_workbook(wb, fpath)
                    return True
            return False

    return with_retry(_do_delete, max_retries=3, retry_interval=1.0)


def replace_mirror(dir_path: str, coaches: list) -> int:
    """用手机备份的 coaches[] 整体替换本地教练档案（只读镜像，以手机为真源）。

    v67（李哥拍板 D2）：教练数据重心在 Android（薪资结算 / 排课 / 团队管理都在那边），
    PC 只有 6 个基础字段且不参与同步 → 同一份数据两端各存一份，早晚对不上。
    改为：PC 教练页只读展示手机备份里的教练，不再提供 PC 端独立编辑。

    - 输入来自 Android BackupManager 导出的 coaches[]（name / phone / specialty / status），
      角色、入职日期、备注三列留空 —— **不动同步协议**，Android 导什么就展示什么
    - status == '在职' → is_active=True；休假 / 离职 → False
    - 幂等：整体替换，同一份备份重复导入结果一致（不留陈旧行）
    - 安全网：恢复流程自带「恢复前自动备份」，覆盖前数据可回滚

    @param coaches: [{'name','phone','specialty','status'}]
    @return 写入的教练条数（姓名为空的行被跳过）
    """
    fpath = ensure_file(dir_path)
    rows = []
    for c in coaches or []:
        if not isinstance(c, dict):
            continue
        name = str(c.get('name') or '').strip()
        if not name:
            continue
        rows.append((
            name,
            str(c.get('phone') or '').strip(),
            str(c.get('specialty') or '').strip(),
            str(c.get('status') or '在职').strip() == '在职',
        ))

    def _do_replace():
        with file_lock(fpath, timeout=5.0):
            wb = load_workbook(fpath)
            ws = wb[COACH_SHEET]
            # 清空数据行（保留表头）—— 整体替换语义，避免手机端删掉的教练在 PC 残留
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            meta = {'coaches': {}}
            for i, (name, phone, specialty, is_active) in enumerate(rows, start=2):
                ws.cell(row=i, column=1, value=name)
                ws.cell(row=i, column=2, value=phone)
                ws.cell(row=i, column=3, value=None)       # 角色：备份未导出
                ws.cell(row=i, column=4, value=specialty)
                ws.cell(row=i, column=5, value=None)       # 入职日期：备份未导出
                ws.cell(row=i, column=6, value=None)       # 备注：备份未导出
                meta['coaches'][name] = {'is_active': is_active, 'updated_at': _now_ms()}
            _write_meta(wb, meta)
            atomic_save_workbook(wb, fpath)
            return len(rows)

    return with_retry(_do_replace, max_retries=3, retry_interval=1.0)
