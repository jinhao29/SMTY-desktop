# -*- coding: utf-8 -*-
"""数据层：课后反馈存储。

反馈记录保存到档案目录下的「课堂反馈.xlsx」。
工作表：反馈记录
列结构：日期 | 学员 | 训练内容 | 完成度(%) | 学员状态 | 教练评语 | 下次课建议

设计：
- 反馈与签到通过「日期+学员」关联
- 完成度记录为百分比整数
- 教练评语支持多行文本
"""
import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from file_lock import file_lock

FEEDBACK_FILE = '课堂反馈.xlsx'
FEEDBACK_SHEET = '反馈记录'

FEEDBACK_HEADERS = ['日期', '学员', '训练内容', '完成度(%)', '学员状态', '教练评语', '下次课建议']

THIN = Side(style='thin', color='888888')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
HEADER_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
HEADER_FONT = Font(bold=True, color='FFFFFF', name='微软雅黑')


def _feedback_file_path(dir_path: str) -> str:
    return os.path.join(dir_path, FEEDBACK_FILE)


def ensure_file(dir_path: str) -> str:
    """确保反馈文件存在。"""
    fpath = _feedback_file_path(dir_path)
    if os.path.exists(fpath):
        return fpath
    os.makedirs(dir_path, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = FEEDBACK_SHEET
    for i, h in enumerate(FEEDBACK_HEADERS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER
    widths = [12, 12, 28, 10, 12, 36, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'
    with file_lock(fpath):
        wb.save(fpath)
    return fpath


def save_feedback(dir_path: str, student_name: str, training_content: str,
                  completion: int, student_state: str, coach_comment: str,
                  next_suggestion: str, date_str: str = '') -> bool:
    """保存一条课后反馈记录。"""
    fpath = ensure_file(dir_path)
    wb = load_workbook(fpath)
    ws = wb[FEEDBACK_SHEET]
    if not date_str:
        date_str = datetime.now().strftime('%Y-%m-%d')
    new_row = ws.max_row + 1 if ws.cell(row=ws.max_row, column=1).value else ws.max_row
    if ws.cell(row=new_row, column=1).value:
        new_row += 1
    values = [date_str, student_name, training_content, completion,
              student_state, coach_comment, next_suggestion]
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=new_row, column=i, value=v)
        cell.border = BORDER
        cell.alignment = CENTER if i <= 5 else LEFT
    with file_lock(fpath):
        wb.save(fpath)
    return True


def get_feedback_history(dir_path: str, student_name: str = None, limit: int = 20) -> list:
    """获取反馈历史记录。"""
    fpath = ensure_file(dir_path)
    wb = load_workbook(fpath)
    ws = wb[FEEDBACK_SHEET]
    result = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        if not name:
            continue
        if student_name and str(name).strip() != student_name:
            continue
        result.append({
            'date': ws.cell(row=r, column=1).value or '',
            'name': str(name).strip(),
            'content': ws.cell(row=r, column=3).value or '',
            'completion': ws.cell(row=r, column=4).value or 0,
            'state': ws.cell(row=r, column=5).value or '',
            'comment': ws.cell(row=r, column=6).value or '',
            'next': ws.cell(row=r, column=7).value or '',
            'row': r,
        })
    result.sort(key=lambda x: str(x['date']), reverse=True)
    return result[:limit]
